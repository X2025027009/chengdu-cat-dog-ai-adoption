"""Build the 10039 technical report from Markdown, including reproducible figures.

Run with python-docx and Pillow available. No network, cloud writes or release edits.
Use export_word_pdf.ps1 on Windows to update fields and export the same DOCX.
"""
from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
CN = "Microsoft YaHei"
BLACK = RGBColor(0, 0, 0)


def figures():
    directory = ROOT / "figures"
    directory.mkdir(exist_ok=True)
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 30)
    small = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 25)

    def base():
        image = Image.new("RGB", (1600, 630), "white")
        return image, ImageDraw.Draw(image)

    def box(draw, xy, lines):
        draw.rectangle(xy, outline="#7b858d", width=2, fill="#f5f7f9")
        x0, y0, x1, y1 = xy
        for index, line in enumerate(lines):
            draw.text(((x0 + x1) / 2, (y0 + y1) / 2 + (index - (len(lines) - 1) / 2) * 44), line, font=font, fill="black", anchor="mm")

    def arrow(draw, start, end):
        draw.line((start, end), fill="#47525c", width=4)
        x, y = end
        draw.polygon([(x, y), (x - 14, y - 8), (x - 14, y + 8)], fill="#47525c")

    image, draw = base()
    labels = [["负责人", "需求与边界"], ["AI", "方案与实现"], ["负责人", "预览与纠偏"], ["离线检查", "本地提交"]]
    for index, lines in enumerate(labels):
        x = 45 + index * 390
        box(draw, (x, 160, x + 320, 345), lines)
        if index < 3:
            arrow(draw, (x + 320, 250), (x + 380, 250))
    draw.text((800, 465), "审阅不通过则返回修改　发布由委托方另行执行", font=font, fill="black", anchor="mm")
    image.save(directory / "fig1-intent-loop.png")

    image, draw = base()
    box(draw, (30, 210, 320, 400), ["七页静态前端", "云端运行路径"])
    destinations = [
        (35, ["公开读与会话", "RLS 约束"], ["PostgreSQL", "公开视图与会话"]),
        (225, ["yard-api", "业务校验与管理"], ["数据库 · 企业微信", "业务写入与群通知"]),
        (415, ["ai-stream", "SSE 与工具循环"], ["数据库 · 模型服务", "检索、工具与响应"]),
    ]
    for y, lines, services in destinations:
        box(draw, (440, y, 960, y + 140), lines)
        arrow(draw, (320, 305), (435, y + 70))
        box(draw, (1080, y, 1580, y + 140), services)
        arrow(draw, (960, y + 70), (1075, y + 70))
    draw.text((800, 595), "数据库为同一业务库；本地审核隔离，不发起上述云端请求", font=small, fill="black", anchor="mm")
    image.save(directory / "fig2-system-boundary.png")

    image, draw = base()
    start, scale = 430, 1.6
    tracks = [("整页水平推移", 0, 360), ("首屏内容独立浮出", 100, 420), ("详情照片弹出", 0, 380), ("详情信息抽屉", 170, 440)]
    for index, (label, delay, duration) in enumerate(tracks):
        y = 65 + index * 105
        draw.text((35, y + 20), label, font=font, fill="black")
        x = start + int(delay * scale)
        draw.rectangle((x, y, x + int(duration * scale), y + 65), fill="#dbe5ed", outline="#7b858d", width=2)
        draw.text((x + 10, y + 18), f"延后 {delay}ms / {duration}ms", font=small, fill="black")
    draw.text((800, 555), "减少动态效果：取消空间动画，保留内容与结果状态", font=font, fill="black", anchor="mm")
    image.save(directory / "fig3-motion-timing.png")


def font(run, size=None, bold=None):
    run.font.name = CN
    run.font.color.rgb = BLACK
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), CN)


def inline(paragraph, text, size=11):
    for part in re.split(r"(`[^`]+`|\*\*.+?\*\*)", text):
        bold = part.startswith("**")
        value = part[2:-2] if bold else part[1:-1] if part.startswith("`") else part
        font(paragraph.add_run(value), size, bold)


def field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    for tag, value in (("begin", None), ("instrText", instruction), ("separate", None), ("t", placeholder), ("end", None)):
        node = OxmlElement(f"w:{'fldChar' if tag in ('begin', 'separate', 'end') else tag}")
        if tag in ("begin", "separate", "end"):
            node.set(qn("w:fldCharType"), tag)
        else:
            node.text = value
        run._r.append(node)


def table(doc, rows):
    count = len(rows[0])
    result = doc.add_table(rows=1, cols=count)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    widths = [5.8, 11.2] if count == 2 else [4.5, 6.25, 6.25]
    if count == 3 and rows[0][0] == "案例":
        widths = [3.0, 8.4, 5.6]
    elif count == 3 and rows[0][0] == "脚本":
        widths = [7.0, 2.8, 7.2]
    for column, width in zip(result.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "5"), ("color", "D9D9D9")):
            node.set(qn(f"w:{key}"), value)
        borders.append(node)
    result._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = result.rows[0] if index == 0 else result.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for column, value in enumerate(values):
            cell = row.cells[column]
            cell.width = Cm(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margin = OxmlElement("w:tcMar")
            for side in ("top", "start", "bottom", "end"):
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), "110")
                node.set(qn("w:type"), "dxa")
                margin.append(node)
            cell._tc.get_or_add_tcPr().append(margin)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            inline(paragraph, value, 10)
            if index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E7ECF0")
                cell._tc.get_or_add_tcPr().append(shading)
                for run in paragraph.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    figures()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    for name in ("TOC 1", "TOC 2"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "TOC 1", "TOC 2"):
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = CN
        style.font.color.rgb = BLACK
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.35
    for name in ("TOC 1", "TOC 2"):
        style = doc.styles[name]
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.0
    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.size = Pt({1: 17, 2: 13, 3: 12}[level])
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(15)
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(30)
    font(title.add_run("成都猫狗小院"), 30, True)
    paragraph = doc.add_paragraph()
    font(paragraph.add_run("AI 协同开发与系统实现技术文档"), 17)
    doc.add_paragraph("2026 第九届传智杯全国大学生数智创新与 AI 应用大赛")
    doc.add_paragraph("AI创新应用挑战赛")
    doc.add_paragraph("2026 年 10 月")
    # The source JPEG carries provenance metadata not recognized by python-docx.
    # Decode to an unchanged-pixel PNG for embedding; never rewrite the original.
    build = ROOT / "build"
    build.mkdir(exist_ok=True)
    cover = build / "cover-team-photo.png"
    with Image.open(ROOT.parent / "assets/team-photo.jpg") as image:
        image.convert("RGB").save(cover)
    doc.add_picture(str(cover), width=Cm(16.0))
    doc.add_page_break()
    paragraph = doc.add_paragraph()
    font(paragraph.add_run("目录"), 20, True)
    field(doc.add_paragraph(), ' TOC \\o "1-2" \\h \\z \\u ', "更新目录")
    doc.add_page_break()
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(footer, " PAGE ", "1")
    lines = (ROOT / "技术文档-正文.md").read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith("# "):
            index += 1
            continue
        if index < 6:
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+", value) for value in values):
                    rows.append(values)
                index += 1
            table(doc, rows)
            continue
        image = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
        if image:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(ROOT / image.group(2)), width=Cm(16.4))
            index += 1
            continue
        heading = re.match(r"^(#{2,3})\s+(.+)$", line)
        if heading:
            paragraph = doc.add_paragraph(style=f"Heading {len(heading.group(1)) - 1}")
            inline(paragraph, heading.group(2), 17 if len(heading.group(1)) == 2 else 13)
            for run in paragraph.runs:
                run.bold = True
        else:
            paragraph = doc.add_paragraph()
            if re.match(r"^[图表]\s*\d+\s", line):
                paragraph.paragraph_format.keep_with_next = True
                inline(paragraph, line, 10)
                for run in paragraph.runs:
                    run.bold = True
            else:
                inline(paragraph, line)
        index += 1
    doc.core_properties.title = "成都猫狗小院技术文档"
    doc.core_properties.subject = "AI创新应用挑战赛 人机协作 架构与复现"
    doc.core_properties.author = "成都猫狗小院"
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings._element.append(update)
    # Bundled default templates may carry a blue Title paragraph border.
    # Remove paragraph borders from styles and direct content, not table borders.
    for border in doc.styles.element.xpath(".//w:pBdr") + doc._element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    target = ROOT / "成都猫狗小院-技术文档.docx"
    doc.save(target)
    print(f"Saved {target}")


if __name__ == "__main__":
    main()
